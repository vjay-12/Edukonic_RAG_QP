import io
from typing import List, Dict, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from app.services.text_utils import (
    format_chemical_formulas,
    strip_option_prefix,
    strip_sub_question_prefix,
    get_roman_numeral,
    format_marks
)

def export_paper_docx(paper_data: Dict[str, Any], questions: List[Dict[str, Any]]) -> io.BytesIO:
    """Generate a cleanly formatted Question Paper DOCX file ready for school printing."""
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    config = paper_data.get("config", {})
    header = config.get("header", {})
    school_name = header.get("school_name", "Edukonic Model Senior Secondary School")
    exam_name = header.get("exam_name", paper_data.get("title", "Science Examination"))
    date_str = header.get("date", "")
    duration = header.get("duration", "1 Hour")
    total_marks = config.get("total_marks", sum(q.get("marks", 0) for q in questions))

    # School Name Heading
    p_school = doc.add_paragraph()
    p_school.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_school = p_school.add_run(school_name.upper())
    r_school.bold = True
    r_school.font.size = Pt(14)
    r_school.font.name = "Calibri"

    # Exam Title
    p_exam = doc.add_paragraph()
    p_exam.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_exam = p_exam.add_run(exam_name)
    r_exam.bold = True
    r_exam.font.size = Pt(12)
    r_exam.font.name = "Calibri"

    # Sub-header: Class, Subject, Time, Max Marks
    class_num = config.get("class", 10)
    subject = config.get("subject", "Science")
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run(f"Class: {class_num}  |  Subject: {subject}  |  Time: {duration}  |  Max. Marks: {total_marks}")
    r_meta.bold = True
    r_meta.font.size = Pt(10.5)

    # Student Identification Line (Requirement 7)
    p_student = doc.add_paragraph()
    p_student.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_student.paragraph_format.space_before = Pt(4)
    p_student.paragraph_format.space_after = Pt(4)
    r_stud = p_student.add_run("Name: __________________________   Roll No: ____________   Date: ____________")
    r_stud.font.size = Pt(10)
    r_stud.font.color.rgb = RGBColor(60, 60, 60)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # General Instructions
    instructions = header.get("instructions", [])
    if instructions:
        p_inst_title = doc.add_paragraph()
        r_it = p_inst_title.add_run("General Instructions:")
        r_it.bold = True
        r_it.font.size = Pt(10)
        p_inst_title.paragraph_format.space_after = Pt(2)

        for ins in instructions:
            p_ins = doc.add_paragraph(style="List Bullet")
            r_ins = p_ins.add_run(ins)
            r_ins.font.size = Pt(9.5)
            p_ins.paragraph_format.space_after = Pt(1)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Group questions by section and sort by template configuration
    sec_map: Dict[str, List[Dict[str, Any]]] = {}
    for q in questions:
        sec = q.get("section", "A")
        if sec not in sec_map:
            sec_map[sec] = []
        sec_map[sec].append(q)

    configured_sections = [s.get("id", "") for s in config.get("sections", []) if isinstance(s, dict)]
    sec_order = sorted(sec_map.keys(), key=lambda s: configured_sections.index(s) if s in configured_sections else 999)

    # Output sections
    for sec in sec_order:
        sec_qs = sec_map[sec]
        m_each = sec_qs[0].get("marks", 1) if sec_qs else 1

        p_sec = doc.add_paragraph()
        p_sec.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_sec = p_sec.add_run(f"SECTION {sec}")
        r_sec.bold = True
        r_sec.font.size = Pt(11)
        r_sub = p_sec.add_run(f"  ({m_each} mark{'s' if m_each > 1 else ''} each)")
        r_sub.italic = True
        r_sub.font.size = Pt(10)
        p_sec.paragraph_format.space_before = Pt(8)
        p_sec.paragraph_format.space_after = Pt(4)

        for q in sec_qs:
            c = q.get("content", {})
            slot = q.get("slot_no", 1)
            raw_stem = c.get("stem", "")
            subs = c.get("sub_questions")
            if subs and isinstance(subs, list) and len(subs) > 0:
                from app.services.text_utils import clean_stem_for_subquestions
                stem = clean_stem_for_subquestions(raw_stem, subs)
            else:
                stem = format_chemical_formulas(raw_stem)
            marks = q.get("marks", 1)

            # Question stem with right marks
            p_q = doc.add_paragraph()
            r_num = p_q.add_run(f"Q{slot}. ")
            r_num.bold = True
            r_num.font.size = Pt(10)

            r_stem = p_q.add_run(stem)
            r_stem.font.size = Pt(10)

            # Marks tag
            r_marks = p_q.add_run(f"  [{marks}]")
            r_marks.bold = True
            r_marks.font.size = Pt(9.5)
            p_q.paragraph_format.space_after = Pt(2)

            # MCQ Options (Strip 'A)', 'B)' prefixes and format formulas)
            options = c.get("options")
            if options and isinstance(options, list):
                labels = ["(A)", "(B)", "(C)", "(D)"]
                for i, raw_opt in enumerate(options):
                    lbl = labels[i] if i < len(labels) else f"({i+1})"
                    cleaned_opt = strip_option_prefix(raw_opt)
                    formatted_opt = format_chemical_formulas(cleaned_opt)

                    p_opt = doc.add_paragraph()
                    p_opt.paragraph_format.left_indent = Inches(0.3)
                    p_opt.paragraph_format.space_after = Pt(1)
                    r_lbl = p_opt.add_run(f"{lbl} ")
                    r_lbl.bold = True
                    r_lbl.font.size = Pt(9.5)
                    r_val = p_opt.add_run(formatted_opt)
                    r_val.font.size = Pt(9.5)

            # Sub-questions (Render consistently with Roman numerals for all types: Case, LA, SA)
            sub_qs = c.get("sub_questions")
            if sub_qs and isinstance(sub_qs, list):
                for i, sq in enumerate(sub_qs):
                    roman = get_roman_numeral(i)
                    raw_sub_stem = sq.get("stem", "")
                    clean_sub = strip_sub_question_prefix(raw_sub_stem)
                    formatted_sub = format_chemical_formulas(clean_sub)

                    p_sub = doc.add_paragraph()
                    p_sub.paragraph_format.left_indent = Inches(0.3)
                    p_sub.paragraph_format.space_after = Pt(2)
                    r_sub_num = p_sub.add_run(f"{roman} ")
                    r_sub_num.bold = True
                    r_sub_num.font.size = Pt(9.5)
                    r_sub_stem = p_sub.add_run(formatted_sub)
                    r_sub_stem.font.size = Pt(9.5)
                    r_sub_m = p_sub.add_run(f"  [{sq.get('marks', 1)}]")
                    r_sub_m.bold = True
                    r_sub_m.font.size = Pt(9)

            p_q.paragraph_format.space_after = Pt(4)

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream

def export_answer_key_docx(paper_data: Dict[str, Any], questions: List[Dict[str, Any]]) -> io.BytesIO:
    """Generate Answer Key and Step-wise Marking Scheme DOCX file."""
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    config = paper_data.get("config", {})
    header = config.get("header", {})
    school_name = header.get("school_name", "Edukonic Model Senior Secondary School")
    exam_name = header.get("exam_name", paper_data.get("title", "Science Examination"))
    total_marks = config.get("total_marks", sum(q.get("marks", 0) for q in questions))

    # Header
    p_school = doc.add_paragraph()
    p_school.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_school = p_school.add_run(f"{school_name.upper()}\nMARKING SCHEME & ANSWER KEY")
    r_school.bold = True
    r_school.font.size = Pt(13)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run(f"{exam_name}  |  Class: {config.get('class', 10)}  |  Subject: {config.get('subject', 'Science')}  |  Max Marks: {total_marks}")
    r_meta.bold = True
    r_meta.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    for q in sorted(questions, key=lambda x: x.get("slot_no", 1)):
        c = q.get("content", {})
        slot = q.get("slot_no", 1)
        marks = q.get("marks", 1)
        sec = q.get("section", "A")
        q_type = q.get("type", "mcq")
        src = c.get("source", {})

        # Canonical Source Reference: "NCERT Class 10 Science, Ch <no> <title>, <section heading>, p.<page> (chunk <id>)"
        citation_str = src.get("citation")
        if not citation_str:
            pages = src.get("pages", [])
            p_str = f"pp. {min(pages)}-{max(pages)}" if len(pages) > 1 else (f"p. {pages[0]}" if pages else "p. 1")
            chunk_id = src.get("primary_chunk_id") or (src.get("chunk_ids") or [1])[0]
            ch_no = q.get("chapter_no", 1)
            ch_title = src.get("chapter_title", "")
            sec_hd = src.get("section_heading", "Textbook Section")
            if src.get("source_type") == "reference_paper":
                citation_str = src.get("reference_name", "CBSE Reference Sample Paper")
            else:
                citation_str = f"NCERT Class 10 Science, Ch {ch_no} {ch_title}, {sec_hd}, {p_str} (chunk {chunk_id})"

        p_q = doc.add_paragraph()
        marks_str = format_marks(marks)
        r_q = p_q.add_run(f"Q{slot} [Section {sec} · {q_type.upper()} · {marks_str}]")
        r_q.bold = True
        r_q.font.size = Pt(10.5)

        corr = c.get("correct_option")
        raw_ans = c.get("answer", "")
        ans = format_chemical_formulas(raw_ans)

        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = Inches(0.2)
        if corr:
            r_corr = p_ans.add_run(f"Correct Option: ({corr}) [1 mark]\n")
            r_corr.bold = True
            r_corr.font.size = Pt(10)

        r_ans_lbl = p_ans.add_run("Model Answer: ")
        r_ans_lbl.bold = True
        r_ans = p_ans.add_run(f"{ans}\n")
        r_ans.font.size = Pt(9.5)

        # Marking Scheme steps with format_marks (1 mark / 2 marks)
        ms = c.get("marking_scheme")
        if ms and isinstance(ms, list):
            r_ms_title = p_ans.add_run("Marking Scheme:\n")
            r_ms_title.bold = True
            r_ms_title.font.size = Pt(9.5)
            for step in ms:
                raw_st = step.get("step", "")
                st = format_chemical_formulas(raw_st)
                st_marks = format_marks(step.get("marks", 1))
                r_step = p_ans.add_run(f"  • {st} — [{st_marks}]\n")
                r_step.font.size = Pt(9)

        # Source citation
        r_src = p_ans.add_run(f"Source Reference: {citation_str}")
        r_src.italic = True
        r_src.font.size = Pt(8.5)

        p_ans.paragraph_format.space_after = Pt(6)

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream
