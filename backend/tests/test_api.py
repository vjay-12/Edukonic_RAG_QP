import io
import pytest
from fastapi.testclient import TestClient
from docx import Document

from app.main import app
from app.config import settings

client = TestClient(app)

def test_api_health():
    res = client.get("/api/health")
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "ok"
    assert data["embed_dim"] == 768

def test_api_auth_login():
    # Bad password
    bad_res = client.post("/api/auth/login", json={"password": "wrong-password"})
    assert bad_res.status_code == 401

    # Good password
    good_res = client.post("/api/auth/login", json={"password": settings.APP_PASSWORD})
    assert good_res.status_code == 200
    token_data = good_res.json()
    assert token_data["authenticated"] is True
    assert "token" in token_data

def test_api_chapters_and_citations():
    res = client.get("/api/chapters")
    assert res.status_code == 200
    chapters = res.json()
    assert len(chapters) >= 3
    ch_nums = [c["chapter_no"] for c in chapters]
    assert 1 in ch_nums
    assert 5 in ch_nums
    assert 11 in ch_nums

    # Test chunk citation
    # Find chunk id from first chapter
    c_res = client.get("/api/chunks/490")
    if c_res.status_code == 200:
        chunk = c_res.json()
        assert "citation" in chunk
        assert "text" in chunk

def test_api_blueprints():
    res = client.get("/api/blueprints")
    assert res.status_code == 200
    bps = res.json()
    assert len(bps) >= 2
    titles = [b.get("title", "") for b in bps]
    assert any("Unit Test" in t for t in titles)

def test_api_papers_history_and_detail():
    # List papers
    res = client.get("/api/papers")
    assert res.status_code == 200
    papers = res.json()
    assert len(papers) >= 1
    p1 = papers[0]
    p_id = p1["id"]

    # Detail
    det_res = client.get(f"/api/papers/{p_id}")
    assert det_res.status_code == 200
    det = det_res.json()
    assert det["id"] == p_id
    assert len(det["questions"]) > 0
    assert det["blueprint_report"] is not None

def test_api_question_patch_and_finalize():
    # Fetch paper 1
    det_res = client.get("/api/papers/1")
    if det_res.status_code == 200:
        det = det_res.json()
        q1 = det["questions"][0]
        qid = q1["id"]

        # Accept question
        patch_res = client.patch(f"/api/papers/1/questions/{qid}", json={"status": "accepted"})
        assert patch_res.status_code == 200

        # Finalize paper
        fin_res = client.post("/api/papers/1/finalize")
        assert fin_res.status_code == 200
        fin_data = fin_res.json()
        assert fin_data["saved_to_bank_count"] >= 1

        # Check bank
        bank_res = client.get("/api/bank")
        assert bank_res.status_code == 200
        bank_items = bank_res.json()
        assert len(bank_items) >= 1

def test_api_export_docx():
    # Find any saved paper from API
    res_list = client.get("/api/papers")
    assert res_list.status_code == 200
    papers = res_list.json()
    assert len(papers) > 0, "At least one saved paper must exist to test export"
    paper_id = papers[0]["id"]

    # 1. Export question paper
    res_paper = client.get(f"/api/papers/{paper_id}/export?format=docx&part=paper")
    assert res_paper.status_code == 200
    assert res_paper.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert f'filename="Paper_{paper_id}.docx"' in res_paper.headers.get("content-disposition", "")
    assert res_paper.content[:4] == b"PK\x03\x04"
    # Assert python-docx opens the file cleanly without corruption
    doc_paper = Document(io.BytesIO(res_paper.content))
    assert len(doc_paper.paragraphs) > 0

    # 2. Export answer key
    res_key = client.get(f"/api/papers/{paper_id}/export?format=docx&part=key")
    assert res_key.status_code == 200
    assert res_key.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert f'filename="Answer_Key_Paper_{paper_id}.docx"' in res_key.headers.get("content-disposition", "")
    assert res_key.content[:4] == b"PK\x03\x04"
    # Assert python-docx opens the key cleanly without corruption
    doc_key = Document(io.BytesIO(res_key.content))
    assert len(doc_key.paragraphs) > 0


def test_api_question_regenerate_mocked(monkeypatch):
    """Verify single question regeneration endpoint with mocked LLM generation."""
    from app.db import SessionLocal
    from app.models import Paper, PaperQuestion

    with SessionLocal() as db:
        test_paper = Paper(
            title="Isolated Test Paper for Regeneration",
            config={"class": 10, "subject": "Science", "total_marks": 2},
            status="ready",
            progress=100
        )
        db.add(test_paper)
        db.commit()
        db.refresh(test_paper)

        test_q = PaperQuestion(
            paper_id=test_paper.id,
            slot_no=1,
            section="A",
            type="vsa",
            marks=2,
            chapter_no=11,
            topic="Initial Topic",
            content={"stem": "Initial stem"},
            status="draft"
        )
        db.add(test_q)
        db.commit()
        db.refresh(test_q)
        test_paper_id = test_paper.id
        test_qid = test_q.id

    mock_content = {
        "stem": "Define electric potential difference and state its SI unit.",
        "options": None,
        "correct_option": None,
        "answer": "Work done per unit charge moving between two points. V = W/Q. SI unit is volt (V).",
        "marking_scheme": [
            {"step": "Work done per unit charge (V = W/Q)", "marks": 1},
            {"step": "SI unit is volt", "marks": 1}
        ],
        "topic": "Electric Potential",
        "concept_key": "potential_difference",
        "source": {"chunk_ids": [561], "pages": [178]}
    }
    mock_val = {"verdict": "pass", "issues": []}
    mock_embs = [[0.1] * 768]
    mock_key = "potential_difference"
    mock_rxns = []
    mock_tags = ["potential_difference"]

    monkeypatch.setattr(
        "app.api.papers.generate_single_slot",
        lambda **kwargs: (mock_content, mock_val, mock_embs, mock_key, mock_rxns, mock_tags)
    )

    try:
        regen_res = client.post(
            f"/api/papers/{test_paper_id}/questions/{test_qid}/regenerate",
            json={"instruction": "Focus on mathematical definition and SI unit."}
        )
        assert regen_res.status_code == 200
        data = regen_res.json()
        assert data["status"] == "ok"
        assert data["question"]["id"] == test_qid
        assert "Define electric potential difference" in data["question"]["content"]["stem"]
    finally:
        with SessionLocal() as db:
            db.query(PaperQuestion).filter(PaperQuestion.paper_id == test_paper_id).delete()
            db.query(Paper).filter(Paper.id == test_paper_id).delete()
            db.commit()


